#!/usr/bin/env python3
"""Keep every table in an R1 DOCX on one page, with its caption.

Pandoc writes tables without pagination hints, so Word may move a table's
last row, or its header row, to the next page. This marks each row as
unsplittable and every paragraph in all but a table's last row as
keep-with-next. Captions that immediately precede a table are kept with it.

Usage: finalize_r1_docx.py FILE.docx [FILE.docx ...]   (edits in place)
"""

from __future__ import annotations

import sys

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_flag(parent, tag: str) -> None:
    """Add an empty on/off property element (e.g. w:keepNext) if missing."""
    if parent.find(qn(tag)) is None:
        parent.append(OxmlElement(tag))


def keep_tables_together(path: str) -> int:
    document = Document(path)
    for table in document.tables:
        rows = table.rows
        for index, row in enumerate(rows):
            _set_flag(row._tr.get_or_add_trPr(), "w:cantSplit")
            if index == len(rows) - 1:
                continue
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _set_flag(paragraph._p.get_or_add_pPr(), "w:keepNext")
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            ppr = previous.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                previous.insert(0, ppr)
            _set_flag(ppr, "w:keepNext")
    document.save(path)
    return len(document.tables)


if __name__ == "__main__":
    for target in sys.argv[1:]:
        print(target, keep_tables_together(target), "tables")
